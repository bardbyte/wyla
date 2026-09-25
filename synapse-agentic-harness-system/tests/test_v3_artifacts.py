"""The charting uplift, as the validator and its helpers hold it:
numbers format themselves from the data (number_format / format_value,
twinned in the surfaces' js/artifacts-render.js on the same cases), a
written heuristic picks the picture (choose_visual, branch by branch),
twelve chart kinds validate strictly with their own encodings and
limits, format overrides ride the spec, a dashboard names its grid,
the artifact tool fills a missing kind and says why, chart_rows builds
the spec the heuristic asked for, and the PPTX export draws every
kind or lands a table — never a crash."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

from sahs.assistant.artifacts import (CHART_KINDS, EXAMPLES, FORMAT_KINDS,
                                      choose_for_series, choose_visual,
                                      format_value, number_format,
                                      validate_artifact)

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from test_v3_dashboards import _kit, compiled  # noqa: E402,F401

CASES = json.loads((HERE / "fixtures" / "number_format_cases.json")
                   .read_text(encoding="utf-8"))
PROV = {"status": "exploratory", "meridian_line": "a look, not a claim"}


# ─── formatting from the data ────────────────────────────────


def test_the_shared_cases_format_as_pinned():
    """The fixture's answers were reviewed by hand; the JS twin runs the
    same file (apps/synapse_admin/tests/test_artifacts_render.py)."""
    assert len(CASES) >= 20
    for case in CASES:
        fmt = number_format(case["values"], case.get("hint"))
        got = {k: fmt[k] for k in case["expect"]}
        assert got == case["expect"], case["name"]
        assert [format_value(v, fmt) for v in case["values"] + [None]] \
            == case["formatted"], case["name"]


def test_decimals_follow_the_data_never_a_fixed_default():
    assert number_format([12.5, 13])["decimals"] == 1
    assert number_format([12.5, 13.25])["decimals"] == 2
    assert number_format([1, 2, 3])["decimals"] == 0
    assert number_format([1.123456])["decimals"] == 4        # the cap
    # a percent as a fraction counts its decimals after the ×100
    pct = number_format([0.125, 0.5], {"name": "rate"})
    assert (pct["kind"], pct["scale"], pct["decimals"]) == ("percent", 100, 1)
    assert format_value(0.125, pct) == "12.5%"
    # given in points, the decimals are the data's own
    pts = number_format([12.5, 13], {"unit": "%"})
    assert (pts["scale"], pts["decimals"]) == (1, 1)
    assert format_value(12.5, pts) == "12.5%"
    # the range decides: a single value past 1 means points
    assert number_format([0.5, 1.5], {"unit": "%"})["scale"] == 1
    assert number_format([1.0, 0.2], {"unit": "%"})["scale"] == 100


def test_metadata_speaks_before_the_values():
    # explicit format beats the unit beats the name beats the values
    assert number_format([1, 2], {"name": "spend", "unit": "%",
                                  "format": "count"})["kind"] == "count"
    assert number_format([1, 2], {"name": "spend", "unit": "%"})["kind"] \
        == "percent"
    assert number_format([1, 2], {"name": "spend"})["kind"] == "currency"
    assert number_format([1, 2])["kind"] == "number"
    fmt = number_format([1234.5678], {"format": {"decimals": 1,
                                                 "unit": "kg"}})
    assert format_value(1234.5678, fmt) == "1,234.6 kg"
    assert number_format([1], {"format": {"kind": "nonsense"}})["kind"] \
        == "number"


def test_compact_units_grouping_and_nulls():
    big = number_format([150000, 2500000, 980000], {"name": "gmv"})
    assert big["compact"] and format_value(2500000, big) == "$2.5M" \
        or format_value(2500000, big) == "2.5M"
    small = number_format([1500, 2500], {"name": "gmv"})
    assert not small["compact"] and format_value(1500, small) == "1,500"
    ids = number_format([150000, 250000], {"name": "merchant_id"})
    assert not ids["compact"] and format_value(150000, ids) == "150000"
    assert format_value(None, big) == "—"
    assert format_value("", big) == "—"
    assert format_value(-1234.5, number_format([-1234.5])) == "-1,234.5"
    assert format_value(-0.0001, number_format([0.5])) == "0.0"
    assert format_value("n/a", big) == "n/a"


# ─── the selection heuristic, branch by branch ───────────────


def _cats(n, measure="spend", label="country"):
    return [{label: f"c{i}", measure: float(i)} for i in range(n)]


def test_one_number_is_a_kpi_with_or_without_a_delta():
    assert choose_visual([{"spend": 5.0}], None)["kind"] == "kpi"
    pick = choose_visual([{"spend": 5.0, "prior_spend": 4.0}], None)
    assert pick["kind"] == "kpi" and pick["compare"] == "prior_spend"
    pick = choose_visual([{"a": 5.0, "b": 4.0}], None, "a vs b")
    assert pick["kind"] == "kpi" and pick["compare"] == "b"
    assert "reason" in pick and pick["reason"].endswith(".")


def test_time_on_x_is_a_line_area_combo_or_small_multiples():
    rows = [{"month": f"2026-0{i}", "spend": 1.0 * i, "orders": 2.0 * i}
            for i in range(1, 7)]
    pick = choose_visual(rows, None)
    assert pick["kind"] == "line" and pick["x"] == "month"
    assert pick["y"] == ["spend", "orders"] and pick["sort"] == "x"
    one = [{"month": r["month"], "spend": r["spend"]} for r in rows]
    assert choose_visual(one, None, "cumulative volume")["kind"] == "area"
    assert choose_visual(one, None)["kind"] == "line"
    plan = [{"month": r["month"], "spend": r["spend"],
             "plan_spend": r["spend"] * 0.9} for r in rows]
    pick = choose_visual(plan, None)
    assert pick["kind"] == "combo" and pick["y"] == ["spend", "plan_spend"]
    wide = [{"month": r["month"], **{f"m{i}": 1.0 for i in range(13)}}
            for r in rows]
    assert choose_visual(wide, None)["kind"] == "small_multiples"
    seven = [{"month": r["month"], **{f"m{i}": 1.0 for i in range(7)}}
             for r in rows]
    assert choose_visual(seven, None)["kind"] == "small_multiples"
    split = [{"month": r["month"], "country": c, "spend": 1.0}
             for r in rows for c in ("US", "CA")]
    pick = choose_visual(split, None)
    assert pick["kind"] == "line" and pick["series"] == "country"
    many = [{"month": r["month"], "country": f"c{i}", "spend": 1.0}
            for r in rows for i in range(13)]
    pick = choose_visual(many, None)
    assert pick["kind"] == "small_multiples" and pick["series"] == "country"


def test_one_category_and_a_measure_is_a_sorted_bar():
    pick = choose_visual(_cats(5), None)
    assert (pick["kind"], pick["x"], pick["y"], pick["sort"]) \
        == ("bar", "country", ["spend"], "-y")
    assert choose_visual(_cats(9), None)["kind"] == "hbar"        # > 8
    long = [{"name": "a very long category label here", "spend": 1.0},
            {"name": "b", "spend": 2.0}]
    assert choose_visual(long, None)["kind"] == "hbar"
    pick = choose_visual(_cats(41), None)
    assert pick["kind"] == "table" and pick["top_n"] == 15 \
        and pick["chart"] == "hbar" and pick["sort"] == "-y"
    grouped = [{"country": f"c{i}", "spend": 1.0, "orders": 2.0}
               for i in range(4)]
    assert choose_visual(grouped, None)["y"] == ["spend", "orders"]


def test_part_of_a_whole_is_a_stack_never_a_pie():
    pick = choose_visual(_cats(4), None, "share of spend by country")
    assert pick["kind"] == "stacked_bar" and pick["series"] == "country"
    pct = [{"country": f"c{i}", "share_pct": 25.0} for i in range(4)]
    assert choose_visual(pct, None)["kind"] == "percent_bar"
    # too many parts to stack: a plain bar
    assert choose_visual(_cats(7), None, "share of spend")["kind"] == "bar"


def test_two_measures_a_distribution_a_matrix_a_bridge():
    pick = choose_visual([{"a": float(i), "b": 2.0 * i} for i in range(30)],
                         None)
    assert (pick["kind"], pick["x"], pick["y"]) == ("scatter", "a", ["b"])
    pick = choose_visual([{"x": float(i)} for i in range(30)], None)
    assert pick["kind"] == "histogram" and pick["y"] == ["x"]
    assert choose_visual(_cats(5), None, "the distribution")["kind"] \
        == "histogram"
    matrix = [{"day": d, "country": c, "n": 1.0}
              for d in ("Mon", "Tue", "Wed") for c in ("US", "CA")]
    pick = choose_visual(matrix, None)
    assert pick["kind"] == "heatmap" and pick["series"] == "day" \
        and pick["x"] == "country"
    pick = choose_visual(_cats(5, measure="delta", label="step"), None,
                         "a bridge from start to end")
    assert pick["kind"] == "waterfall"


def test_nothing_to_draw_is_a_table():
    assert choose_visual([], None)["kind"] == "table"
    assert choose_visual([{"t": "dw.x"}, {"t": "dw.y"}], None)["kind"] \
        == "table"
    # an identifier column with one measure: still a bar by the label
    ids = [{"merchant_id": f"m{i}", "spend": 1.0} for i in range(3)]
    assert choose_visual(ids, None)["kind"] == "bar"


def test_the_series_form_of_the_heuristic():
    assert choose_for_series(EXAMPLES["line"]["series"])["kind"] == "combo"
    assert choose_for_series(EXAMPLES["bar"]["series"])["kind"] == "bar"
    assert choose_for_series(EXAMPLES["area"]["series"])["kind"] == "line"
    assert choose_for_series([])["kind"] == "table"


# ─── every kind validates, strictly ──────────────────────────


def test_every_kind_has_an_example_that_validates():
    assert set(EXAMPLES) == set(CHART_KINDS) and len(CHART_KINDS) == 12
    for kind, example in EXAMPLES.items():
        out, problems = validate_artifact("chart", example, build_id="b")
        assert problems == [], (kind, problems)
        assert out["kind"] == kind
        assert out["watermark"] == "EXPLORATORY"
    assert validate_artifact("chart", EXAMPLES["line"], build_id="b")[0][
        "reference"] == {"value": 950000, "label": "Plan"}
    binned = validate_artifact("chart", EXAMPLES["histogram"],
                               build_id="b")[0]
    assert binned["binned"] and binned["bins"] == 5
    assert len(binned["series"][0]["points"]) == 5
    assert sum(p[1] for p in binned["series"][0]["points"]) == 20
    assert "–" in binned["series"][0]["points"][0][0]


def test_kind_aliases_and_unknown_kinds():
    for alias, kind in (("horizontal_bar", "hbar"), ("stacked", "stacked_bar"),
                        ("100%", "percent_bar"), ("facets", "small_multiples"),
                        ("bridge", "waterfall"), ("column", "bar")):
        spec = dict(EXAMPLES[kind], kind=alias)
        out, problems = validate_artifact("chart", spec, build_id="b")
        assert problems == [] and out["kind"] == kind, alias
    _, problems = validate_artifact("chart", {"kind": "pie", "series": [],
                                              "provenance": PROV})
    assert problems[0]["code"] == "chart_kind"
    assert "leave kind out" in problems[0]["hint"]


def _refused(spec, code):
    out, problems = validate_artifact("chart", dict(spec, provenance=PROV))
    assert out is None and code in [p["code"] for p in problems], problems
    return problems


def test_each_kind_requires_its_own_encoding():
    _refused({"kind": "heatmap", "x": ["a"], "y": ["b"], "values": [[1, 2]]},
             "heatmap_matrix")
    _refused({"kind": "heatmap", "x": ["a"], "y": ["b"], "values": [["x"]]},
             "heatmap_matrix")
    _refused({"kind": "heatmap", "x": ["a"], "y": ["b"], "values": [[None]]},
             "heatmap_empty")
    _refused({"kind": "heatmap", "x": [str(i) for i in range(61)],
              "y": ["b"], "values": [[1] * 61]}, "heatmap_size")
    _refused({"kind": "histogram", "values": [1]}, "histogram_values")
    _refused({"kind": "scatter", "series": [{"name": "s", "points": [
        ["a", 1]]}]}, "scatter_x")
    _refused({"kind": "combo", "series": [
        {"name": "a", "role": "bar", "points": [["x", 1]]},
        {"name": "b", "role": "bar", "points": [["x", 1]]}]}, "combo_roles")
    _refused({"kind": "waterfall", "series": [
        {"name": "a", "points": [["x", 1]]},
        {"name": "b", "points": [["x", 1]]}]}, "waterfall_series")
    _refused({"kind": "line", "series": [
        {"name": f"s{i}", "points": [["x", 1]]} for i in range(7)]},
        "chart_series_many")
    _refused({"kind": "small_multiples", "series": [
        {"name": f"s{i}", "points": [["x", 1]]} for i in range(25)]},
        "chart_series_many")
    ok, problems = validate_artifact("chart", {
        "kind": "small_multiples", "provenance": PROV, "series": [
            {"name": f"s{i}", "points": [["x", 1]]} for i in range(24)]})
    assert problems == [] and len(ok["series"]) == 24
    _refused({"kind": "line", "series": [{"name": "s", "points": [
        [str(i), 1] for i in range(2001)]}]}, "chart_points_many")
    _refused({"kind": "bar", "series": [{"name": "s", "points": [["a", 1]]}],
              "reference": "high"}, "reference_line")
    _refused({"kind": "bar", "series": [{"name": "s", "points": [["a", 1]]}],
              "sort": "sideways"}, "sort_unknown")
    _refused({"kind": "bar", "series": [{"name": "s", "points": [["a", 1]]}],
              "format": {"kind": "nope"}}, "format_unknown")
    # a combo with roles unstated: first bar, then lines
    out, problems = validate_artifact("chart", {
        "kind": "combo", "provenance": PROV, "series": [
            {"name": "a", "points": [["x", 1]]},
            {"name": "b", "points": [["x", 1]]}]})
    assert problems == [] and [s["role"] for s in out["series"]] \
        == ["bar", "line"]


def test_format_overrides_ride_the_spec():
    table, problems = validate_artifact("table", {
        "columns": [{"key": "r", "label": "Rate", "format": "percent"},
                    {"key": "s", "label": "Spend", "unit": "USD",
                     "format": {"decimals": 0, "compact": False}}],
        "rows": [{"r": 0.125, "s": 1234.5}], "provenance": PROV})
    assert problems == []
    assert table["columns"][0]["format"] == {"kind": "percent"}
    assert table["columns"][1] == {"key": "s", "label": "Spend",
                                   "unit": "USD",
                                   "format": {"decimals": 0,
                                              "compact": False}}
    _, problems = validate_artifact("table", {
        "columns": [{"key": "r", "label": "R", "format": "bogus"}],
        "rows": [{"r": 1}], "provenance": PROV})
    assert [p["code"] for p in problems] == ["format_unknown"]
    kpi, problems = validate_artifact("kpi", {
        "value": 1234567.8, "unit": "USD", "delta": -0.042,
        "delta_format": "percent", "compare_label": "vs last month",
        "provenance": PROV})
    assert problems == []
    assert kpi["delta_format"] == {"kind": "percent"}
    assert kpi["compare_label"] == "vs last month"
    chart, problems = validate_artifact("chart", {
        "kind": "line", "format": "currency", "x_format": {"kind": "date"},
        "reason": "because", "sort": "-y", "provenance": PROV,
        "series": [{"name": "s", "unit": "EUR", "format": "count",
                    "points": [["a", 1]]}]})
    assert problems == []
    assert chart["format"] == {"kind": "currency"}
    assert chart["x_format"] == {"kind": "date"}
    assert chart["series"][0]["unit"] == "EUR"
    assert chart["series"][0]["format"] == {"kind": "count"}
    assert chart["reason"] == "because" and chart["sort"] == "-y"
    assert set(FORMAT_KINDS) >= {"percent", "currency", "count", "ratio",
                                 "duration", "date"}


def test_a_dashboard_names_its_grid_and_each_panel_its_span():
    spec = {"grid": 3, "panels": [
        {"type": "kpi", "spec": {"value": 1.0, "provenance": PROV}},
        {"type": "chart", "span": 2, "spec": EXAMPLES["heatmap"]},
        {"type": "chart", "span": 9, "spec": EXAMPLES["bar"]}]}
    out, problems = validate_artifact("dashboard", spec, build_id="b")
    assert problems == []
    assert out["grid"] == 3
    assert "span" not in out["panels"][0]
    assert out["panels"][1]["span"] == 2
    assert "span" not in out["panels"][2]          # 9 is not 1..3
    _, problems = validate_artifact("dashboard", dict(spec, grid=4))
    assert [p["code"] for p in problems] == ["dashboard_grid"]
    out, problems = validate_artifact("dashboard", {
        "grid": 1, "panels": [{"type": "chart", "span": 3,
                               "spec": EXAMPLES["bar"]}]})
    assert problems == [] and out["panels"][0]["span"] == 1


# ─── the tools use the heuristic ─────────────────────────────


def test_the_artifact_tool_fills_a_missing_kind_and_says_why(
        compiled, tmp_path):
    tools, _state, store, sid, _ws = _kit(compiled, tmp_path)
    naked = {"series": EXAMPLES["bar"]["series"], "provenance": PROV}
    made = tools["artifact"].fn("chart", "By country", json.dumps(naked))
    assert made["ok"], made
    spec = store.get_artifact(made["artifact_id"])["spec"]
    assert spec["kind"] == "bar" and spec["sort"] == "-y"
    assert spec["reason"].startswith("Spend by x: bars sorted by value")
    # the model's own reason and kind are kept; an alias is normalized
    own = dict(naked, kind="horizontal bar", reason="the labels are long")
    made = tools["artifact"].fn("chart", "By country", json.dumps(own))
    assert made["ok"]
    spec = store.get_artifact(made["artifact_id"])["spec"]
    assert spec["kind"] == "hbar" and spec["reason"] == "the labels are long"
    # one point in one series is a kpi, not a chart: refused, teaching
    one = {"series": [{"name": "spend", "points": [["total", 5.0]]}],
           "provenance": PROV}
    made = tools["artifact"].fn("chart", "Total", json.dumps(one))
    assert made["error"] == "artifact refused"
    assert made["problems"][0]["code"] == "chart_kind"
    assert "kpi" in made["problems"][0]["detail"]


def test_chart_rows_builds_the_spec_the_heuristic_asked_for():
    from sahs.assistant.loop import _chart_spec
    rows = [{"day": "Mon", "country": c, "n": float(i + 1)}
            for i, c in enumerate(("US", "CA"))] + [
        {"day": "Tue", "country": c, "n": float(i + 3)}
        for i, c in enumerate(("US", "CA"))]
    pick = choose_visual(rows, None)
    assert pick["kind"] == "heatmap"
    spec = _chart_spec(rows, "heatmap", pick["x"], pick["y"],
                       pick["series"], pick, pick["reason"])
    assert spec["x"] == ["US", "CA"] and spec["y"] == ["Mon", "Tue"]
    assert spec["values"] == [[1.0, 2.0], [3.0, 4.0]]
    out, problems = validate_artifact("chart", dict(spec, provenance=PROV))
    assert problems == [] and out["reason"] == pick["reason"]
    # a long list: the top N as a horizontal bar, sorted
    many = [{"country": f"c{i:02d}", "spend": float(i)} for i in range(50)]
    pick = choose_visual(many, None)
    spec = _chart_spec(many, "hbar", pick["x"], pick["y"], "", pick,
                       pick["reason"])
    assert len(spec["series"][0]["points"]) == 15
    assert spec["series"][0]["points"][0] == ["c49", 49.0]
    # a split column: one series per value, capped at six by total
    split = [{"month": f"2026-0{m}", "country": f"c{i}", "spend": float(i)}
             for m in range(1, 4) for i in range(8)]
    pick = choose_visual(split, None)
    spec = _chart_spec(split, "line", pick["x"], pick["y"], pick["series"],
                       pick, pick["reason"])
    assert [s["name"] for s in spec["series"]] == [f"c{i}" for i in
                                                    range(7, 1, -1)]
    # a distribution: raw values, binned by the validator
    dist = [{"x": float(i)} for i in range(30)]
    pick = choose_visual(dist, None)
    spec = _chart_spec(dist, "histogram", "", pick["y"], "", pick, "")
    out, problems = validate_artifact("chart", dict(spec, provenance=PROV))
    assert problems == [] and out["binned"]
    # a scatter: numeric pairs
    pairs = [{"a": float(i), "b": 2.0 * i} for i in range(30)]
    pick = choose_visual(pairs, None)
    spec = _chart_spec(pairs, "scatter", pick["x"], pick["y"], "", pick, "")
    assert spec["series"][0]["points"][1] == [1.0, 2.0]
    assert validate_artifact("chart", dict(spec, provenance=PROV))[1] == []


# ─── the export draws every kind or lands a table ────────────


def test_the_export_draws_every_kind_or_lands_a_table():
    pytest.importorskip("pptx")
    from pptx import Presentation
    from sahs.assistant.export import artifact_pptx
    for kind, example in EXAMPLES.items():
        spec, problems = validate_artifact("chart", example, build_id="b")
        assert problems == []
        deck = artifact_pptx({"type": "chart", "title": kind, "version": 1,
                              "spec": spec})
        prs = Presentation(__import__("io").BytesIO(deck))
        slide = prs.slides[1]
        charts = [s for s in slide.shapes if s.has_chart]
        tables = [s for s in slide.shapes if s.has_table]
        notes = slide.notes_slide.notes_text_frame.text
        if kind in ("heatmap", "waterfall"):
            assert tables and not charts, kind
            assert "no native PowerPoint chart" in notes
        else:
            assert charts and not tables, kind
        assert "Why this chart" in notes or "reason" not in spec
    # a table exports its numbers formatted, a kpi its delta
    table, _ = validate_artifact("table", {
        "columns": [{"key": "spend", "label": "Spend", "unit": "USD"},
                    {"key": "rate", "label": "Rate"}],
        "rows": [{"spend": 1234.5, "rate": 0.125}], "provenance": PROV},
        build_id="b")
    prs = Presentation(__import__("io").BytesIO(artifact_pptx(
        {"type": "table", "title": "t", "version": 1, "spec": table})))
    cells = [c.text for c in prs.slides[1].shapes[1].table.iter_cells()]
    assert cells == ["Spend", "Rate", "$1,234.5", "12.5%"]
    # a broken spec never crashes the export
    deck = artifact_pptx({"type": "chart", "title": "g", "version": 1,
                          "spec": {"kind": "combo", "series": [
                              {"name": "a", "points": []}]}})
    assert deck[:2] == b"PK"
