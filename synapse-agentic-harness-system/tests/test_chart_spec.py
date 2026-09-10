"""A chart's spec tells the truth about its axis: every series rides
one shared x axis, a null is a gap (a forecast is null over the
actuals and starts at the last actual), a dashed flag marks a
projection, and the pptx export draws the same axis."""

from __future__ import annotations

from sahs.assistant.artifacts import validate_artifact
from sahs.assistant.skills_loader import builtin_skills

PROV = {"status": "exploratory", "meridian_line": "a look, not a claim"}


def test_nulls_are_gaps_and_dashed_marks_a_forecast():
    spec = {"kind": "line", "unit": "USD", "provenance": PROV,
            "series": [
                {"name": "Actuals", "points": [["2026-01", 802], ["2026-02", 850],
                                               ["2026-03", 900], ["2026-04", None],
                                               ["2026-05", None], ["2026-06", None]]},
                {"name": "Forecast", "dashed": True,
                 "points": [["2026-03", 900], ["2026-04", 930], ["2026-05", 960],
                            ["2026-06", 1020]]}]}
    out, problems = validate_artifact("chart", spec, build_id="b1")
    assert problems == [], problems
    assert out["series"][0]["points"][3] == ["2026-04", None]
    assert out["series"][1]["dashed"] is True
    assert "dashed" not in out["series"][0]
    # a series of nothing but nulls is no series
    _, problems = validate_artifact("chart", {
        "kind": "line", "provenance": PROV,
        "series": [{"name": "x", "points": [["a", None], ["b", None]]}]},
        build_id="b1")
    assert [p["code"] for p in problems] == ["chart_points"]
    assert "null for a gap" in problems[0]["hint"]
    # a string y is still refused
    _, problems = validate_artifact("chart", {
        "kind": "bar", "provenance": PROV,
        "series": [{"name": "x", "points": [["a", "12"]]}]}, build_id="b1")
    assert [p["code"] for p in problems] == ["chart_points"]
    # seven series is too many to read
    _, problems = validate_artifact("chart", {
        "kind": "line", "provenance": PROV,
        "series": [{"name": f"s{i}", "points": [["a", i]]} for i in range(7)]},
        build_id="b1")
    assert "chart_series_many" in [p["code"] for p in problems]


def test_the_export_draws_one_axis_for_every_series():
    import pytest
    pytest.importorskip("pptx")
    from sahs.assistant.export import _chart_slide
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _chart_slide(slide, {"kind": "line", "series": [
        {"name": "Actuals", "points": [["2026-01", 1], ["2026-02", 2]]},
        {"name": "Forecast", "points": [["2026-02", 2], ["2026-03", 3]]}]})
    chart = [s for s in slide.shapes if s.has_chart][0].chart
    cats = list(chart.plots[0].categories)
    assert cats == ["2026-01", "2026-02", "2026-03"]
    values = [list(s.values) for s in chart.plots[0].series]
    assert values[0] == [1, 2, None] and values[1] == [None, 2, 3]


def test_the_charts_skill_is_built_in_and_says_the_rules():
    pack = {p.name: p for p in builtin_skills()}["charts"]
    assert pack.origin == "built-in"
    for piece in ("ONE shared x axis", "null", "dashed: true",
                  "never a pie", "two y axes", "Bars start at zero",
                  "starts at\n   the last actual point"):
        assert piece.lower() in pack.text.lower(), piece
